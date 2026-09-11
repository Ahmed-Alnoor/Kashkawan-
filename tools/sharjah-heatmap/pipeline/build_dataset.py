#!/usr/bin/env python3
"""Turn the two Sharjah prospecting workbooks into one dataset for the heatmap.

    python3 build_dataset.py <community_split.xlsx> <d11_target_list.xlsx> -o ../app/data.js

Output is a JS file that assigns `window.SHARJAH_DATA`, so the viewer can load
it with a plain <script> tag and stay a static, dependency-free page.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from gazetteer import ALIASES, AREAS, LANDMARKS  # noqa: E402

import openpyxl  # noqa: E402

# --------------------------------------------------------------------------
# Taxonomy: both workbooks use their own category vocabularies. Everything is
# folded into one `category` list, and categories roll up into 8 `sector`s so
# the map only ever needs eight distinguishable colours.
# --------------------------------------------------------------------------
CATEGORY_MAP = {
    # workbook A - "Activity Group"
    "Healthcare": "Healthcare & Clinics",
    "Beauty & Wellness": "Beauty, Spa & Wellness",
    "Fitness & Sports": "Fitness & Sports",
    "Food & Beverage": "Food & Beverage",
    "Premium Retail": "Retail & Department Stores",
    "Luxury & Jewellery": "Luxury & Jewellery",
    "Home & Lifestyle": "Home & Lifestyle",
    "Finance & Banking": "Banking & Financial Services",
    "Insurance": "Insurance",
    "Legal": "Legal",
    "Accounting, Audit & Tax": "Accounting, Audit & Tax",
    "Business & Management Consulting": "Consulting & Business Services",
    "Engineering & Architecture": "Engineering & Architecture",
    "Real Estate & Property": "Real Estate & Property",
    "Technology & Electronics": "Technology & Electronics",
    "Technology": "Technology & Electronics",
    "Marketing & Media": "Marketing, Media & Creative",
    "Education & Training": "Education & Training",
    "Education & Lifestyle": "Education & Training",
    # workbook B - "Category"
    "Medical Centre & Polyclinic": "Healthcare & Clinics",
    "Dental Clinic": "Dental",
    "Aesthetic, Derm & Cosmetic Clinic": "Aesthetic & Cosmetic",
    "Physio, Rehab & Wellness": "Physio, Rehab & Wellness",
    "Veterinary": "Veterinary",
    "Pharmacy & Medical Supplies": "Pharmacy & Medical Supplies",
    "Spa, Salon & Personal Care": "Beauty, Spa & Wellness",
    "Fitness & Sports Facility": "Fitness & Sports",
    "Retail, F&B & Hospitality": "Retail, F&B & Hospitality",
    "Financial Services & Insurance": "Banking & Financial Services",
    "Professional & Business Services": "Consulting & Business Services",
    "Construction & Contracting": "Construction & Contracting",
    "Heavy & Industrial Manufacturing": "Manufacturing & Industrial",
    "Trading, Wholesale & Distribution": "Trading & Distribution",
    "Logistics, Shipping & Transport": "Logistics & Transport",
    "Technology, Media & Marketing": "Technology, Media & Marketing",
    "Energy, Oil, Gas & Utilities": "Energy & Utilities",
    "Automotive & Workshops": "Automotive & Workshops",
    "Institution / Government": "Institution / Government",
    "Other / Unclassified": "Other / Unclassified",
}

SECTOR_MAP = {
    "Healthcare & Clinics": "Healthcare & Wellness",
    "Dental": "Healthcare & Wellness",
    "Aesthetic & Cosmetic": "Healthcare & Wellness",
    "Physio, Rehab & Wellness": "Healthcare & Wellness",
    "Veterinary": "Healthcare & Wellness",
    "Pharmacy & Medical Supplies": "Healthcare & Wellness",
    "Beauty, Spa & Wellness": "Healthcare & Wellness",
    "Fitness & Sports": "Healthcare & Wellness",
    "Banking & Financial Services": "Professional & Financial",
    "Insurance": "Professional & Financial",
    "Legal": "Professional & Financial",
    "Accounting, Audit & Tax": "Professional & Financial",
    "Consulting & Business Services": "Professional & Financial",
    "Real Estate & Property": "Professional & Financial",
    "Retail & Department Stores": "Retail & Hospitality",
    "Luxury & Jewellery": "Retail & Hospitality",
    "Home & Lifestyle": "Retail & Hospitality",
    "Food & Beverage": "Retail & Hospitality",
    "Retail, F&B & Hospitality": "Retail & Hospitality",
    "Manufacturing & Industrial": "Industrial & Manufacturing",
    "Energy & Utilities": "Industrial & Manufacturing",
    "Construction & Contracting": "Construction & Engineering",
    "Engineering & Architecture": "Construction & Engineering",
    "Trading & Distribution": "Logistics & Trade",
    "Logistics & Transport": "Logistics & Trade",
    "Automotive & Workshops": "Logistics & Trade",
    "Technology & Electronics": "Technology & Education",
    "Technology, Media & Marketing": "Technology & Education",
    "Marketing, Media & Creative": "Technology & Education",
    "Education & Training": "Technology & Education",
    "Institution / Government": "Public & Other",
    "Other / Unclassified": "Public & Other",
}

TIER_MAP = {
    "A - Industrial & Large Business": "A · Industrial & Large",
    "B - Clinic / Aesthetic / Wellness": "B · Clinic & Wellness",
    "C - Corporate & Professional": "C · Corporate & Professional",
    "D - Retail, F&B & Other": "D · Retail, F&B & Other",
    "E - Institution / Government": "E · Institution & Government",
}

PRIORITY_MAP = {
    "P1 - Core target": "P1 · Core target",
    "P2 - High": "P2 · High",
    "P3 - Secondary": "P3 · Secondary",
    "P4 - Watchlist": "P4 · Watchlist",
    "Not a target": "Not a target",
    # workbook A grades
    "A": "P1 · Core target",
    "B": "P2 · High",
    "C": "P3 · Secondary",
}

UNKNOWN_AREA = "Area unknown"

# Vehicle workshops, garages and the auto-parts trade are out of scope for this
# list. Matched three ways because the source workbooks tag them inconsistently:
# by broad category, by detailed activity, and by what the business calls itself.
DROP_CATEGORIES = {"Automotive & Workshops"}
DROP_DETAIL = re.compile(
    r"^(auto\s|automotive\b|auto/|auto-)|"
    r"(auto\s*(spare\s*)?parts|auto\s*repair|auto\s*accessories|auto\s*dealership|"
    r"automotive\s*services|car\s*(rental|wash|care|service)|tyre|garage)",
    re.I)
DROP_NAME = re.compile(r"\b(workshop|garage|auto\s*(garage|repair|maintenance|service|care))\b", re.I)

# A business is "free zone" if it came off the free-zone register or sits in a
# free-zone district. Kept in the dataset, excluded from the default view.
FREEZONE_AREAS = {
    "Hamriyah Free Zone", "Sharjah Airport Free Zone", "Sharjah Free Zone",
    "Sharjah Media City", "Sharjah Publishing City",
    "Sharjah Research Technology and Innovation Park",
}


def is_excluded(rec: dict) -> str | None:
    """Return the reason this record is dropped, or None to keep it."""
    if rec.get("category") in DROP_CATEGORIES:
        return "automotive category"
    if DROP_DETAIL.search(rec.get("detail") or ""):
        return "automotive activity"
    if DROP_NAME.search(rec.get("name") or ""):
        return "workshop / garage"
    return None

COORD_RE = re.compile(r"([-+]?\d{1,2}\.\d{3,}),\s*([-+]?\d{1,3}\.\d{3,})")


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------
def clean(value) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).replace(" ", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return "" if text.lower() in {"none", "nan", "n/a", "-", "not verified"} else text


def norm_key(text: str) -> str:
    text = unicodedata.normalize("NFKD", text.lower())
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def stable_unit(seed: str, salt: str) -> float:
    """Deterministic float in [0, 1) — same input always gives the same point."""
    digest = hashlib.sha256(f"{seed}|{salt}".encode()).digest()
    return int.from_bytes(digest[:8], "big") / 2**64


def scatter(lat: float, lng: float, radius_m: float, seed: str) -> tuple[float, float]:
    """Place a point inside a disc of `radius_m` around the centroid.

    sqrt on the radial term keeps the points area-uniform rather than bunched
    at the centre, which is what makes a district read as a district.
    """
    angle = stable_unit(seed, "angle") * 2 * math.pi
    dist = math.sqrt(stable_unit(seed, "radius")) * radius_m
    dlat = (dist * math.cos(angle)) / 111_320.0
    dlng = (dist * math.sin(angle)) / (111_320.0 * math.cos(math.radians(lat)))
    return round(lat + dlat, 6), round(lng + dlng, 6)


def resolve_area(raw: str) -> str | None:
    """Map a free-text area string onto a gazetteer key."""
    if not raw:
        return None
    text = raw.strip()
    if text.lower().startswith("area unknown") or text.lower() in {"sharjah", "sharjah, uae"}:
        return None
    if text in AREAS:
        return text
    key = norm_key(text)
    if key in ALIASES:
        return ALIASES[key]
    lookup = {norm_key(k): k for k in AREAS}
    if key in lookup:
        return lookup[key]
    # "Industrial Area 12, Sharjah" and friends
    m = re.search(r"industrial area\s*(\d{1,2})", key)
    if m:
        candidate = f"Industrial Area {int(m.group(1))}"
        if candidate in AREAS:
            return candidate
    for alias_key, canonical in ALIASES.items():
        if alias_key and alias_key in key:
            return canonical
    # longest containment wins so "al majaz 3" beats "al majaz"
    best = None
    for norm, canonical in sorted(lookup.items(), key=lambda kv: -len(kv[0])):
        if len(norm) >= 5 and norm in key:
            best = canonical
            break
    return best


def resolve_landmark(premises: str) -> tuple | None:
    if not premises:
        return None
    key = norm_key(premises)
    for name, entry in LANDMARKS.items():
        if norm_key(name) in key:
            return entry
    return None


def sheet_records(ws, header_row: int) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= header_row:
        return []
    header = [clean(h) for h in rows[header_row]]
    out = []
    for row in rows[header_row + 1:]:
        if not row:
            continue
        record = {h: row[i] if i < len(row) else None for i, h in enumerate(header) if h}
        if clean(record.get("Business Name")):
            out.append(record)
    return out


# --------------------------------------------------------------------------
# per-workbook readers
# --------------------------------------------------------------------------
def read_community_workbook(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = sheet_records(wb["Master List"], 0)
    out = []
    for r in rows:
        group = clean(r.get("Activity Group"))
        category = CATEGORY_MAP.get(group, group or "Other / Unclassified")
        out.append({
            "name": clean(r.get("Business Name")),
            "source": "Premium communities",
            "category": category,
            "detail": clean(r.get("Detailed Activity")) or group,
            "tier": "Premium destination tenant"
            if clean(r.get("Prospect Type")).startswith("Premium")
            else "Core professional / healthcare",
            "priority": PRIORITY_MAP.get(clean(r.get("Priority")), clean(r.get("Priority"))),
            "cluster": clean(r.get("Location Cluster")),
            "community": clean(r.get("Community")),
            "areaRaw": clean(r.get("Area")),
            "premises": clean(r.get("Building / Premises")),
            "unit": clean(r.get("Floor / Unit")),
            "address": clean(r.get("Full Address")),
            "website": clean(r.get("Website")),
            "maps": clean(r.get("Google Maps Search")),
            "verification": clean(r.get("Verification Level")),
            "notes": clean(r.get("Notes")),
            "freezone": False,
        })
    wb.close()
    return out


def read_target_workbook(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    # The healthcare checkpoint sheet carries verified areas, addresses and
    # phone numbers for a subset of rows — fold it in by business name.
    enrich: dict[str, dict] = {}
    if "Healthcare Checkpoint" in wb.sheetnames:
        for r in sheet_records(wb["Healthcare Checkpoint"], 3):
            enrich[norm_key(clean(r.get("Business Name")))] = {
                "area": clean(r.get("Sharjah Area")),
                "address": clean(r.get("Exact Address")),
                "phone": clean(r.get("Phone")),
                "website": clean(r.get("Website")),
                "instagram": clean(r.get("Instagram")),
            }

    out = []
    for sheet, is_freezone in (("Master List", False), ("Freezones", True)):
        if sheet not in wb.sheetnames:
            continue
        for r in sheet_records(wb[sheet], 3):
            name = clean(r.get("Business Name"))
            raw_category = clean(r.get("Category"))
            category = CATEGORY_MAP.get(raw_category, raw_category or "Other / Unclassified")
            area_raw = clean(r.get("Sharjah Area"))
            address = clean(r.get("Best Sharjah Address / Area"))
            website = clean(r.get("Website"))
            instagram = clean(r.get("Instagram"))
            phone = ""
            extra = enrich.get(norm_key(name))
            if extra:
                area_raw = area_raw if not area_raw.startswith(UNKNOWN_AREA) else extra["area"] or area_raw
                address = extra["address"] or address
                phone = extra["phone"]
                website = website or extra["website"]
                instagram = instagram or extra["instagram"]
            out.append({
                "name": name,
                "source": "Free zone register" if is_freezone else "District 11 target list",
                "category": category,
                "detail": clean(r.get("Detailed Category")) or clean(r.get("Sub-sector")) or raw_category,
                "tier": TIER_MAP.get(clean(r.get("Target Tier")), clean(r.get("Target Tier"))),
                "priority": PRIORITY_MAP.get(clean(r.get("Priority")), clean(r.get("Priority"))),
                "cluster": "",
                "community": "",
                "areaRaw": area_raw,
                "premises": "",
                "unit": "",
                "address": address,
                "website": website,
                "instagram": instagram,
                "phone": phone,
                "maps": clean(r.get("Google Maps Link")),
                "verification": clean(r.get("Map Verification Level")),
                "confidence": clean(r.get("Location confidence")),
                "notes": clean(r.get("Location Notes")),
                "freezone": is_freezone,
            })
    wb.close()
    return out


# --------------------------------------------------------------------------
# geocoding
# --------------------------------------------------------------------------
def geocode(rec: dict) -> dict:
    """Attach lat/lng plus an honest precision label.

    exact     – a real coordinate was published in the source
    premises  – a named building/mall with a known point location
    area      – district centroid, scattered inside the district footprint
    none      – Sharjah confirmed but no district; deliberately left off the map
    """
    seed = f"{rec['name']}|{rec['source']}"

    m = COORD_RE.search(rec.get("maps", "") or "")
    if m:
        lat, lng = float(m.group(1)), float(m.group(2))
        if 22 < lat < 27 and 51 < lng < 57:
            rec.update(lat=round(lat, 6), lng=round(lng, 6), precision="exact")
            rec["area"] = resolve_area(rec.get("areaRaw", "")) or rec.get("areaRaw") or "Sharjah"
            rec["areaType"] = AREAS.get(rec["area"], (0, 0, 0, "Mixed urban"))[3]
            return rec

    landmark = resolve_landmark(rec.get("premises", ""))
    if landmark:
        lat, lng, radius, area_name = landmark
        rec["lat"], rec["lng"] = scatter(lat, lng, radius, seed)
        rec["precision"] = "premises"
        rec["area"] = resolve_area(rec.get("areaRaw", "")) or area_name
        rec["areaType"] = AREAS.get(rec["area"], (0, 0, 0, "Mixed urban"))[3]
        return rec

    area = (
        resolve_area(rec.get("areaRaw", ""))
        or resolve_area(rec.get("community", ""))
        or resolve_area(rec.get("cluster", ""))
        or resolve_area(rec.get("address", ""))
    )
    if area and area in AREAS:
        lat, lng, radius, kind = AREAS[area]
        rec["lat"], rec["lng"] = scatter(lat, lng, radius, seed)
        rec["precision"] = "area"
        rec["area"] = area
        rec["areaType"] = kind
        return rec

    rec["lat"] = rec["lng"] = None
    rec["precision"] = "none"
    rec["area"] = UNKNOWN_AREA
    rec["areaType"] = "Not yet traced"
    return rec


# --------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("community_xlsx")
    ap.add_argument("target_xlsx")
    ap.add_argument("-o", "--out", default="data.js")
    args = ap.parse_args()

    records = read_community_workbook(Path(args.community_xlsx))
    records += read_target_workbook(Path(args.target_xlsx))

    dropped = Counter()
    kept_records = []
    for rec in records:
        reason = is_excluded(rec)
        if reason:
            dropped[reason] += 1
        else:
            kept_records.append(rec)
    records = kept_records

    # De-duplicate on name + area. The premium-community workbook is the
    # richer record, so it wins when the same business appears in both.
    by_key: dict[str, dict] = {}
    duplicates = 0
    for rec in records:
        key = f"{norm_key(rec['name'])}|{norm_key(rec.get('areaRaw', ''))}"
        if key in by_key:
            duplicates += 1
            kept = by_key[key]
            for field in ("website", "instagram", "phone", "address", "maps", "notes"):
                if not kept.get(field) and rec.get(field):
                    kept[field] = rec[field]
            continue
        by_key[key] = rec
    records = list(by_key.values())

    # A business listed once against a district and once city-wide is still one
    # business — but a retail brand with a store in two malls is two records,
    # so only the district-less copy is dropped, never a second located one.
    def located(rec: dict) -> bool:
        return bool(resolve_area(rec.get("areaRaw", "")) or resolve_area(rec.get("community", ""))
                    or resolve_area(rec.get("premises", "")) or resolve_landmark(rec.get("premises", "")))

    have_located: set[str] = {norm_key(r["name"]) for r in records if located(r)}
    kept, seen_unlocated = [], set()
    for rec in records:
        key = norm_key(rec["name"])
        if located(rec):
            kept.append(rec)
            continue
        if key in have_located or key in seen_unlocated:
            duplicates += 1
            continue
        seen_unlocated.add(key)
        kept.append(rec)
    records = kept

    for i, rec in enumerate(records):
        geocode(rec)
        rec["sector"] = SECTOR_MAP.get(rec["category"], "Public & Other")
        if rec["area"] in FREEZONE_AREAS:
            rec["freezone"] = True
        rec["id"] = i + 1

    records.sort(key=lambda r: (r["sector"], r["category"], r["name"]))
    for i, rec in enumerate(records):
        rec["id"] = i + 1

    # Compact wire format: short keys, arrays instead of objects, shared
    # dictionaries for the repeated strings. Keeps the payload small enough
    # to ship as one static file.
    dicts = {}

    def dict_index(name: str, value: str) -> int:
        table = dicts.setdefault(name, {})
        if value not in table:
            table[value] = len(table)
        return table[value]

    rows = []
    for rec in records:
        rows.append([
            rec["id"],
            rec["name"],
            dict_index("sector", rec["sector"]),
            dict_index("category", rec["category"]),
            dict_index("tier", rec["tier"] or "Unclassified"),
            dict_index("priority", rec["priority"] or "Unrated"),
            dict_index("area", rec["area"]),
            dict_index("source", rec["source"]),
            dict_index("precision", rec["precision"]),
            rec["lat"],
            rec["lng"],
            rec.get("detail", ""),
            rec.get("address", ""),
            rec.get("website", ""),
            rec.get("maps", ""),
            rec.get("phone", ""),
            rec.get("instagram", ""),
            dict_index("verification", rec.get("verification") or "Not recorded"),
            rec.get("premises", ""),
            dict_index("areaType", rec.get("areaType") or "Not yet traced"),
            1 if rec.get("freezone") else 0,
        ])

    tables = {name: [k for k, _ in sorted(t.items(), key=lambda kv: kv[1])] for name, t in dicts.items()}

    area_points = {
        name: {"lat": v[0], "lng": v[1], "r": v[2], "type": v[3]}
        for name, v in AREAS.items()
    }

    payload = {
        "meta": {
            "generated": "",
            "total": len(records),
            "mapped": sum(1 for r in records if r["lat"] is not None),
            "duplicatesMerged": duplicates,
            "sources": [
                Path(args.community_xlsx).name,
                Path(args.target_xlsx).name,
            ],
        },
        "fields": [
            "id", "name", "sector", "category", "tier", "priority", "area", "source",
            "precision", "lat", "lng", "detail", "address", "website", "maps",
            "phone", "instagram", "verification", "premises", "areaType", "freezone",
        ],
        "tables": tables,
        "areas": area_points,
        "rows": rows,
    }

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    body = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    out_path.write_text(f"window.SHARJAH_DATA={body};\n", encoding="utf-8")

    # --- report -----------------------------------------------------------
    prec = Counter(r["precision"] for r in records)
    print(f"records            {len(records)}")
    print(f"duplicates merged  {duplicates}")
    print(f"free zone          {sum(1 for r in records if r.get('freezone'))} (kept, filtered out by default)")
    if dropped:
        print("excluded:")
        for k, v in dropped.most_common():
            print(f"  {v:6d}  {k}")
    print(f"output             {out_path}  ({out_path.stat().st_size/1024:.0f} KB)")
    print("\nprecision:")
    for k, v in prec.most_common():
        print(f"  {v:6d}  {k}")
    print("\ntop mapped areas:")
    areas = Counter(r["area"] for r in records if r["lat"] is not None)
    for k, v in areas.most_common(15):
        print(f"  {v:6d}  {k}")
    print("\nsectors:")
    for k, v in Counter(r["sector"] for r in records).most_common():
        print(f"  {v:6d}  {k}")
    unmatched = Counter(
        r["areaRaw"] for r in records
        if r["precision"] == "none" and r.get("areaRaw") and not r["areaRaw"].startswith(UNKNOWN_AREA)
    )
    if unmatched:
        print("\nAREA STRINGS THAT DID NOT MATCH THE GAZETTEER:")
        for k, v in unmatched.most_common(25):
            print(f"  {v:6d}  {k}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
